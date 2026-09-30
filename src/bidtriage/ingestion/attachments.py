"""Attachment text extraction (SPEC-01 F5). OCR is optional (extra `ocr`)."""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass

LARGE_PAGES = 40
LARGE_BYTES = 50 * 1024 * 1024
OVERSIZE_BYTES = 200 * 1024 * 1024
ZIP_MAX_FILES = 25
ZIP_MAX_TOTAL = 200 * 1024 * 1024
_UNSAFE = re.compile(r"[\\/\x00-\x1f]")


@dataclass
class ExtractedText:
    text: str | None
    ocr: bool = False
    large_document: bool = False
    error: str | None = None
    pages: int | None = None


def sanitize_filename(name: str) -> str:
    return _UNSAFE.sub("_", name).strip() or "attachment"


def sniff_mime(data: bytes, declared: str) -> str:
    head = data[:8]
    if head.startswith(b"%PDF"):
        return "application/pdf"
    if head.startswith(b"PK\x03\x04"):
        return declared if declared.endswith(("document", "sheet", "zip")) else "application/zip"
    if head.lstrip().lower().startswith((b"<!doctype html", b"<html")):
        return "text/html"
    return declared


def extract_pdf_text(data: bytes, *, ocr: bool = False) -> ExtractedText:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                return ExtractedText(None, error="encrypted")
        pages = len(reader.pages)
    except PdfReadError:
        return ExtractedText(None, error="unreadable")
    if pages > LARGE_PAGES or len(data) > LARGE_BYTES:
        return ExtractedText(None, large_document=True, pages=pages)
    chunks = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        chunks.append(f"[page {i}]\n{t.strip()}")
    text = "\n\n".join(chunks)
    if len(re.sub(r"\W", "", text)) < 20 * max(pages, 1) and ocr:
        return _ocr_pdf(data, pages)
    if len(re.sub(r"\W", "", text)) < 20 * max(pages, 1):
        return ExtractedText(text or None, error="no_text_layer", pages=pages)
    return ExtractedText(text, pages=pages)


def _ocr_pdf(data: bytes, pages: int) -> ExtractedText:
    try:
        import pytesseract
        from pdf2image import convert_from_bytes
    except ImportError:
        return ExtractedText(None, error="ocr_unavailable", pages=pages)
    images = convert_from_bytes(data, dpi=200, first_page=1, last_page=min(pages, 20))
    text = "\n\n".join(
        f"[page {i}]\n{pytesseract.image_to_string(img)}" for i, img in enumerate(images, start=1)
    )
    return ExtractedText(text, ocr=True, pages=pages)


def extract_docx_text(data: bytes) -> ExtractedText:
    try:
        import docx
    except ImportError:
        return ExtractedText(None, error="docx_unavailable")
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception:
        return ExtractedText(None, error="unreadable")
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            parts.append("\t".join(c.text.strip() for c in row.cells))
    return ExtractedText("\n".join(parts))


def list_zip(data: bytes) -> list[tuple[str, bytes]]:
    out: list[tuple[str, bytes]] = []
    total = 0
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for info in z.infolist():
            if info.is_dir() or len(out) >= ZIP_MAX_FILES:
                continue
            if info.filename.lower().endswith(".zip"):
                continue  # nested zips listed by caller, not opened
            if total + info.file_size > ZIP_MAX_TOTAL:
                break
            total += info.file_size
            out.append((sanitize_filename(info.filename.rsplit("/", 1)[-1]), z.read(info)))
    return out


def extract_text(filename: str, mime: str, data: bytes, *, ocr: bool = False) -> ExtractedText:
    if len(data) > OVERSIZE_BYTES:
        return ExtractedText(None, error="oversize")
    mime = sniff_mime(data, mime)
    lower = filename.lower()
    if mime == "application/pdf" or lower.endswith(".pdf") and mime != "text/html":
        return extract_pdf_text(data, ocr=ocr)
    if lower.endswith(".docx"):
        return extract_docx_text(data)
    if mime.startswith("text/plain") or lower.endswith(".txt"):
        return ExtractedText(data.decode("utf-8", errors="replace"))
    return ExtractedText(None)
